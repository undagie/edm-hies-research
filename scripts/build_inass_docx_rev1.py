"""Build revised INASS DOCX from manuscript markdown; colour revisions red vs submit.

Pipeline: pandoc -> style remap -> two-column layout -> template compliance ->
revision colouring (run structure preserved so OMML survives).

Usage: python scripts/build_inass_docx_rev1.py
Output: paper/edm_hies_paper_inass_rev1.docx
"""

from __future__ import annotations

import difflib
import re
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "paper" / "edm_hies_paper_inass.md"
REF = ROOT / "paper" / "edm_hies_paper_inass_submit.docx"
# Official IJIES template (Ver. 2025.5.18) — header/footer adopted verbatim.
TEMPLATE = ROOT / "paper" / "ijies_format_2025.docx"
OUT = ROOT / "paper" / "edm_hies_paper_inass_rev1.docx"

RED = "FF0000"


def build_pandoc(tmp_out: Path) -> None:
    # Manual section numbers in heading text; auto-numbering stripped later.
    work = tmp_out.with_suffix(".md")
    work.write_text(SRC.read_text(encoding="utf-8"), encoding="utf-8")
    cmd = [
        "pandoc",
        str(work),
        "--from=markdown+tex_math_dollars+pipe_tables+raw_html",
        "--to=docx",
        f"--reference-doc={REF}",
        f"--resource-path={ROOT / 'paper'}",
        "-o",
        str(tmp_out),
    ]
    print("running:", " ".join(cmd))
    subprocess.run(cmd, check=True)


def post_process_styles(docx_path: Path) -> None:
    """Remap pandoc's default style names to INASS template names."""
    import zipfile

    with zipfile.ZipFile(docx_path, "r") as zin:
        contents = {name: zin.read(name) for name in zin.namelist()}

    doc = contents["word/document.xml"].decode("utf-8")
    doc = doc.replace('w:val="Heading4"', 'w:val="__H3"')
    doc = doc.replace('w:val="Heading3"', 'w:val="__H2"')
    doc = doc.replace('w:val="Heading2"', 'w:val="__H1"')
    doc = doc.replace('w:val="Heading1"', 'w:val="__TITLE"')
    doc = doc.replace('w:val="__TITLE"', 'w:val="Title"')
    doc = doc.replace('w:val="__H1"', 'w:val="Heading1"')
    doc = doc.replace('w:val="__H2"', 'w:val="Heading2"')
    doc = doc.replace('w:val="__H3"', 'w:val="Heading3"')
    doc = doc.replace('w:val="BodyText"', 'w:val="TextBody"')
    doc = doc.replace('w:val="FirstParagraph"', 'w:val="TextBody"')
    doc = doc.replace('w:val="ImageCaption"', 'w:val="Caption"')

    ref_idx = doc.find(">References<")
    if ref_idx > 0:
        head, tail = doc[:ref_idx], doc[ref_idx:]

        def _flip(match: re.Match) -> str:
            block = match.group(0)
            if re.search(r"<w:t[^>]*>\s*\[\d+\]", block[:2500]):
                return block.replace('w:val="TextBody"', 'w:val="Bibliography"', 1)
            return block

        tail = re.sub(r"<w:p(?:\s[^>]*)?>(?:(?!</w:p>).)*?</w:p>", _flip, tail, flags=re.DOTALL)
        doc = head + tail

    # Strip heading auto-numbering so manual numbers are not doubled.
    styles = contents["word/styles.xml"].decode("utf-8")

    def _strip_style_numbering(match: re.Match) -> str:
        return re.sub(r"<w:numPr>.*?</w:numPr>", "", match.group(0), flags=re.DOTALL)

    styles = re.sub(
        r'<w:style [^>]*w:styleId="Heading[1-9]"[^>]*>.*?</w:style>',
        _strip_style_numbering,
        styles,
        flags=re.DOTALL,
    )

    # Force non-math fonts to Times; italicize code styles (no monospace).
    styles = re.sub(
        r'(<w:rFonts[^>]*w:ascii=")(?!Cambria Math)[^"]*("[^>]*w:hAnsi=")[^"]*(")',
        r"\1Times New Roman\2Times New Roman\3",
        styles,
    )
    for sid in ("VerbatimChar", "SourceCode"):
        m = re.search(rf'(<w:style [^>]*w:styleId="{sid}"[^>]*>)(.*?)(</w:style>)', styles, re.DOTALL)
        if m and "<w:i/>" not in m.group(2):
            rpr = re.search(r"<w:rPr>(.*?)</w:rPr>", m.group(2), re.DOTALL)
            if rpr:
                new_rpr = "<w:rPr>" + rpr.group(1) + "<w:i/>" + "</w:rPr>"
                styles = styles.replace(m.group(0), m.group(1) + m.group(2).replace(rpr.group(0), new_rpr) + m.group(3))
    contents["word/styles.xml"] = styles.encode("utf-8")

    if TEMPLATE.exists():
        with zipfile.ZipFile(TEMPLATE, "r") as ztpl:
            tpl_names = set(ztpl.namelist())
            for part in ("word/header1.xml", "word/footer1.xml"):
                if part in tpl_names and part in contents:
                    contents[part] = ztpl.read(part)

    contents["word/document.xml"] = doc.encode("utf-8")
    with zipfile.ZipFile(docx_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in contents.items():
            zout.writestr(name, data)


def post_process_layout(docx_path: Path) -> None:
    """Single-column front matter, two-column body; wide floats span both columns."""
    import zipfile

    with zipfile.ZipFile(docx_path, "r") as zin:
        contents = {name: zin.read(name) for name in zin.namelist()}
    doc = contents["word/document.xml"].decode("utf-8")

    # Continuous break after Keywords ends the single-column front matter.
    kw_idx = doc.find("Keywords:")
    if kw_idx > 0:
        para_end = doc.find("</w:p>", kw_idx)
        if para_end > 0:
            break_para = (
                "<w:p><w:pPr><w:sectPr>"
                '<w:type w:val="continuous"/>' + _SECT_COMMON + '<w:cols w:space="708"/>'
                "</w:sectPr></w:pPr></w:p>"
            )
            insert_at = para_end + len("</w:p>")
            doc = doc[:insert_at] + break_para + doc[insert_at:]

    tail_sect = doc.rfind("<w:sectPr")
    if tail_sect > 0:
        tail_end = doc.find("</w:sectPr>", tail_sect)
        sect = doc[tail_sect:tail_end]
        sect = re.sub(r"<w:cols[^/>]*/>", "", sect)
        sect = re.sub(r"<w:pgMar[^/>]*/>", _PGMAR, sect)
        sect = re.sub(
            r"(<w:sectPr[^>]*>)",
            r'\1<w:type w:val="continuous"/><w:cols w:num="2" w:space="432"/>',
            sect,
            count=1,
        )
        sect = re.sub(r"<w:type w:val=\"[a-zA-Z]+\"/>(?=(?:(?!<w:type).)*<w:type)", "", sect, flags=re.DOTALL)
        doc = doc[:tail_sect] + sect + doc[tail_end:]

    doc = _wrap_wide_floats(doc)

    contents["word/document.xml"] = doc.encode("utf-8")
    with zipfile.ZipFile(docx_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in contents.items():
            zout.writestr(name, data)


# Page geometry from official IJIES_Format.docx (A4, two-col body).
_PGMAR = (
    '<w:pgMar w:top="1440" w:right="1080" w:bottom="1440" w:left="1080" '
    'w:header="850" w:footer="567" w:gutter="0"/>'
)
_SECT_COMMON = '<w:pgSz w:w="11906" w:h="16838" w:code="9"/>' + _PGMAR
_BREAK_TWO_COL = (
    "<w:p><w:pPr><w:sectPr>"
    '<w:type w:val="continuous"/>' + _SECT_COMMON + '<w:cols w:num="2" w:space="432"/>'
    "</w:sectPr></w:pPr></w:p>"
)
_BREAK_ONE_COL = (
    "<w:p><w:pPr><w:sectPr>"
    '<w:type w:val="continuous"/>' + _SECT_COMMON + '<w:cols w:space="708"/>'
    "</w:sectPr></w:pPr></w:p>"
)

# Notation table stays in-column; other tables wrap as full-width floats.
_NARROW_TABLE_CAPTIONS = ("Table 10.",)


def _wrap_wide_floats(doc: str) -> str:
    body_start = doc.find("<w:body>") + len("<w:body>")
    body_end = doc.rfind("<w:sectPr")
    head, body, tail = doc[:body_start], doc[body_start:body_end], doc[body_end:]

    token_re = re.compile(
        r"<w:tbl>.*?</w:tbl>|<w:p(?:\s[^>]*)?>.*?</w:p>|<w:p(?:\s[^>]*)?/>|<w:[a-zA-Z]+(?:\s[^>]*)?/>|<w:sdt>.*?</w:sdt>",
        flags=re.DOTALL,
    )
    elements: list[str] = []
    pos = 0
    while pos < len(body):
        ws = re.match(r"\s+", body[pos:])
        if ws:
            if elements:
                elements[-1] += ws.group(0)
            pos += ws.end()
            continue
        match = token_re.match(body, pos)
        if match is None:
            print(f"[warn] body tokenisation stalled at offset {pos}: {body[pos:pos+80]!r}; skipping wide-float wrapping")
            return doc
        elements.append(match.group(0))
        pos = match.end()

    def _text_of(el: str) -> str:
        return re.sub(r"<[^>]+>", "", el)

    first_break = next((i for i, el in enumerate(elements) if "<w:sectPr>" in el), 0)

    out: list[str] = []
    i = 0
    while i < len(elements):
        el = elements[i]
        if i <= first_break:
            out.append(el)
            i += 1
            continue
        text = _text_of(el).strip()
        is_figure = "<w:drawing" in el and not el.startswith("<w:tbl>")
        is_table_caption = (
            text.startswith("Table ")
            and not any(text.startswith(n) for n in _NARROW_TABLE_CAPTIONS)
            and i + 1 < len(elements)
            and elements[i + 1].startswith("<w:tbl>")
        )
        if is_figure:
            group = [el]
            if i + 1 < len(elements) and 'w:val="Caption"' in elements[i + 1]:
                group.append(elements[i + 1])
            out.extend([_BREAK_TWO_COL, *group, _BREAK_ONE_COL])
            i += len(group)
        elif is_table_caption:
            group = [el, elements[i + 1]]
            out.extend([_BREAK_TWO_COL, *group, _BREAK_ONE_COL])
            i += 2
        else:
            out.append(el)
            i += 1
    return head + "".join(out) + tail


def post_process_template_compliance(docx_path: Path) -> None:
    """Apply IJIES font/size/alignment specs from the official template."""
    import docx as docx_lib
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    doc = docx_lib.Document(str(docx_path))

    style_specs = {
        "Title": dict(size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER),
        "Heading1": dict(size=12, bold=True),
        "Heading2": dict(size=11, bold=True),
        "Heading3": dict(size=10, bold=True),
        "TextBody": dict(size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=Pt(16.5)),
        "Caption": dict(size=10, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.CENTER),
        "Bibliography": dict(size=11),
    }
    for style in doc.styles:
        spec = style_specs.get(style.style_id)
        if spec is None:
            continue
        if hasattr(style, "font"):
            style.font.size = Pt(spec["size"])
            style.font.name = "Times New Roman"
            if "bold" in spec:
                style.font.bold = spec["bold"]
            if "italic" in spec:
                style.font.italic = spec["italic"]
        if hasattr(style, "paragraph_format"):
            if "align" in spec:
                style.paragraph_format.alignment = spec["align"]
            if "first_line" in spec:
                style.paragraph_format.first_line_indent = spec["first_line"]

    def _set_border(paragraph, edge: str) -> None:
        p_pr = paragraph._p.get_or_add_pPr()
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is None:
            p_bdr = p_pr.makeelement(qn("w:pBdr"), {})
            p_pr.append(p_bdr)
        line = p_bdr.makeelement(qn(f"w:{edge}"), {
            qn("w:val"): "single", qn("w:sz"): "6",
            qn("w:space"): "4", qn("w:color"): "auto",
        })
        p_bdr.append(line)

    body_paras = list(_iter_paragraphs(doc))
    abstract_idx = next(
        (i for i, p in enumerate(body_paras) if p.text.strip().startswith("Abstract:")), None
    )
    for i, para in enumerate(body_paras):
        text = para.text.strip()
        is_front = abstract_idx is not None and i < abstract_idx
        if is_front and para.style.style_id != "Title" and text:
            para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.first_line_indent = Pt(0)
            if text.startswith("* Corresponding"):
                for run in para.runs:
                    run.font.size = Pt(10.5)
        if text.startswith("Abstract:") or text.startswith("Keywords:"):
            para.paragraph_format.first_line_indent = Pt(0)
            for run in para.runs:
                run.font.size = Pt(10)
            _set_border(para, "top" if text.startswith("Abstract:") else "bottom")
        if re.match(r"^Table \d+\.", text):
            para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.first_line_indent = Pt(0)
            for run in para.runs:
                run.font.size = Pt(10)
                run.font.bold = False
        if "<m:oMath" in para._p.xml and len(text) < 5:
            para.paragraph_format.left_indent = Mm(5)
            para.paragraph_format.first_line_indent = Pt(0)

    def _tables_of(container):
        from docx.table import Table

        for child in container.iter_inner_content():
            if isinstance(child, Table):
                yield child

    for table in _tables_of(doc):
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.first_line_indent = Pt(0)
                    for run in para.runs:
                        run.font.size = Pt(10)
        _style_table_booktabs(table)

    doc.save(str(docx_path))


# ECMA-376 child order — inserts must stay schema-valid or Word "repairs".
_TBLPR_ORDER = [
    "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
    "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd",
    "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
    "tblCaption", "tblDescription",
]
_TCPR_ORDER = [
    "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd",
    "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark",
]


def _insert_in_order(parent, child, order: list[str]) -> None:
    from docx.oxml.ns import qn

    tag = child.tag.split("}")[-1]
    idx = order.index(tag)
    after = set(order[idx + 1:])
    for existing in parent:
        if existing.tag.split("}")[-1] in after:
            existing.addprevious(child)
            return
    parent.append(child)


def _rule(edge: str, present: bool):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    el = OxmlElement(f"w:{edge}")
    if present:
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
    else:
        el.set(qn("w:val"), "nil")
    return el


def _style_table_booktabs(table) -> None:
    """Horizontal rules only (top, under header, bottom); no vertical/shading."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            for tag in ("w:tcBorders", "w:shd"):
                for e in tcPr.findall(qn(tag)):
                    tcPr.remove(e)

    tblPr = table._tbl.tblPr
    for e in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(e)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom"):
        borders.append(_rule(edge, True))
    for edge in ("left", "right", "insideH", "insideV"):
        borders.append(_rule(edge, False))
    _insert_in_order(tblPr, borders, _TBLPR_ORDER)

    if table.rows:
        for cell in table.rows[0].cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcb = OxmlElement("w:tcBorders")
            tcb.append(_rule("bottom", True))
            _insert_in_order(tcPr, tcb, _TCPR_ORDER)


def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _iter_paragraphs(container):
    """Yield every paragraph in a python-docx container, including tables."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in container.iter_inner_content():
        if isinstance(child, Paragraph):
            yield child
        elif isinstance(child, Table):
            for row in child.rows:
                for cell in row.cells:
                    yield from _iter_paragraphs(cell)


def _colour_run_red(run) -> None:
    from docx.shared import RGBColor

    run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)


def _changed_char_ranges(old: str, new: str) -> list[tuple[int, int]]:
    """Character ranges in ``new`` that differ from ``old`` (word-level diff)."""
    old_words = old.split()
    new_words = new.split()
    sm = difflib.SequenceMatcher(a=old_words, b=new_words, autojunk=False)
    spans: list[tuple[int, int]] = []
    idx = 0
    starts: list[tuple[int, int]] = []
    for w in new_words:
        pos = new.find(w, idx)
        starts.append((pos, pos + len(w)))
        idx = pos + len(w)
    for tag, _i1, _i2, j1, j2 in sm.get_opcodes():
        if tag in ("replace", "insert") and j2 > j1:
            spans.append((starts[j1][0], starts[j2 - 1][1]))
    return spans


def _alpha_words(text: str) -> set:
    """Letters-only word bag (ignores digits/superscripts/punctuation)."""
    out = set()
    for w in text.split():
        a = "".join(c for c in w.lower() if c.isalpha())
        if a:
            out.add(a)
    return out


def _frontmatter_words(doc) -> set:
    """Submitted front-matter word bag (through Abstract) for fuzzy equality."""
    words = set()
    for p in _iter_paragraphs(doc):
        t = p.text.strip()
        if t.startswith("Abstract:"):
            break
        words |= _alpha_words(t)
    return words


def colour_revisions(docx_path: Path, old_docx: Path) -> dict:
    import docx as docx_lib

    old_doc = docx_lib.Document(str(old_docx))
    old_paras = [_norm(p.text) for p in _iter_paragraphs(old_doc) if _norm(p.text)]
    old_set = set(old_paras)
    old_front_words = _frontmatter_words(old_doc)

    new_doc = docx_lib.Document(str(docx_path))
    seen_abstract = False
    stats = {"unchanged": 0, "partial": 0, "new": 0, "skipped": 0}
    for para in _iter_paragraphs(new_doc):
        text = para.text
        norm = _norm(text)
        if not norm:
            stats["skipped"] += 1
            continue
        if norm.startswith("Abstract:"):
            seen_abstract = True
        if norm in old_set:
            stats["unchanged"] += 1
            continue
        # Submitted headings lacked manual numbers; compare title text only.
        if para.style.style_id.startswith("Heading"):
            unnumbered = _norm(re.sub(r"^\d+(?:\.\d+)*\.?\s+", "", norm))
            if unnumbered in old_set:
                stats["unchanged"] += 1
                continue
        # Front matter re-splits / cosmetic punctuation stay black if words match.
        if para.style.style_id == "Title" or not seen_abstract:
            para_words = _alpha_words(norm)
            if para_words and para_words <= old_front_words:
                stats["unchanged"] += 1
                continue
        close = difflib.get_close_matches(norm, old_paras, n=1, cutoff=0.60)
        if not close:
            for run in para.runs:
                _colour_run_red(run)
            stats["new"] += 1
            continue
        ranges = _changed_char_ranges(close[0], text)
        if not ranges:
            stats["unchanged"] += 1
            continue
        pos = 0
        for run in para.runs:
            run_start, run_end = pos, pos + len(run.text)
            pos = run_end
            if any(run_start < e and run_end > s for s, e in ranges):
                _colour_run_red(run)
        stats["partial"] += 1
    new_doc.save(str(docx_path))
    return stats


def main() -> None:
    if not SRC.exists() or not REF.exists():
        sys.exit(f"Missing {SRC} or {REF}")
    leftover = re.findall(r"%%[A-Z_]+%%", SRC.read_text(encoding="utf-8"))
    if leftover:
        sys.exit(f"Unresolved placeholders in manuscript: {sorted(set(leftover))}")
    tmp = OUT.with_suffix(".building.docx")
    if tmp.exists():
        tmp.unlink()
    build_pandoc(tmp)
    post_process_styles(tmp)
    post_process_layout(tmp)
    post_process_template_compliance(tmp)
    stats = colour_revisions(tmp, REF)
    if OUT.exists():
        OUT.unlink()
    tmp.rename(OUT)
    print(f"Wrote {OUT} ({OUT.stat().st_size:,} bytes)")
    print(f"Revision colouring: {stats}")


if __name__ == "__main__":
    main()
